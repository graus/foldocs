"""Reading a Nextcloud share over WebDAV, against a fake server; then a sync and a build from it."""

import io
import logging
import xml.etree.ElementTree as ET
from dataclasses import replace
from urllib.parse import quote, unquote, urlparse

import docx
import pytest
import requests
from PIL import Image

from foldocs import cache, config, render, sync
from foldocs.formats import DOCX
from foldocs.sources import TooLarge
from foldocs.sources.nextcloud import Nextcloud, parse_share, warned

SHARE = "https://cloud.example.org/s/Tok3n"
DAV = "/public.php/webdav"


def word(text: str) -> bytes:
    d = docx.Document()
    d.add_paragraph("A title", style="Title")
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# path -> properties; a folder has no content type. Times as Nextcloud sends them.
TREE = {
    "/": {},
    "/1 Work": {"fileid": "10"},
    "/1 Work/A post.docx": {
        "fileid": "11",
        "type": DOCX,
        "modified": "Tue, 20 Oct 2026 08:00:00 GMT",
        "creation_time": "0",
        "upload_time": "1767225600",  # 2026-01-01
    },
    "/1 Work/notes.odt": {"fileid": "12", "type": "application/vnd.oasis.opendocument.text"},
    "/1 Work/Sub folder": {"fileid": "13"},
    "/1 Work/Sub folder/Second.docx": {
        "fileid": "14",
        "type": DOCX,
        "modified": "Wed, 21 Oct 2026 09:30:00 GMT",
        "creation_time": "1767312000",  # 2026-01-02
    },
}
FILES = {"/1 Work/A post.docx": word("First words."), "/1 Work/Sub folder/Second.docx": word("More words.")}


def response(path: str, props: dict) -> str:
    folder = "type" not in props
    href = quote(DAV + path.rstrip("/") + ("/" if folder else ""))
    found = [
        f"<d:resourcetype>{'<d:collection/>' if folder else ''}</d:resourcetype>",
        *([f"<oc:fileid>{props['fileid']}</oc:fileid>"] if "fileid" in props else []),
        *([f"<d:getcontenttype>{props['type']}</d:getcontenttype>"] if not folder else []),
        f"<d:getlastmodified>{props.get('modified', 'Mon, 19 Oct 2026 10:00:00 GMT')}</d:getlastmodified>",
        *(f"<nc:{k}>{props[k]}</nc:{k}>" for k in ("creation_time", "upload_time") if k in props),
    ]
    return (
        f"<d:response><d:href>{href}</d:href>"
        f"<d:propstat><d:prop>{''.join(found)}</d:prop><d:status>HTTP/1.1 200 OK</d:status></d:propstat>"
        # What a server answers for a property it doesn't have: empty, and not found.
        "<d:propstat><d:prop><nc:upload_time/></d:prop><d:status>HTTP/1.1 404 Not Found</d:status></d:propstat>"
        "</d:response>"
    )


class Reply:
    def __init__(self, status: int, content: bytes = b""):
        self.status_code, self.content, self.ok = status, content, status < 400

    def raise_for_status(self):
        if not self.ok:
            raise requests.HTTPError(f"{self.status_code} Client Error")

    def iter_content(self, size):
        return iter([self.content])

    def __enter__(self):
        return self

    def __exit__(self, *_):
        pass


class FakeServer:
    def __init__(self, tree=TREE, files=FILES):
        self.tree, self.files, self.requests = tree, files, []

    def _path(self, url: str) -> str:
        path = unquote(urlparse(url).path)
        assert path.startswith(DAV)
        return path[len(DAV) :].rstrip("/") or "/"

    def request(self, method, url, data=None, headers=None, timeout=None, stream=False):
        assert method == "PROPFIND" and headers["Depth"] == "1"
        path = self._path(url)
        self.requests.append(path)
        if path not in self.tree:
            return Reply(404)
        prefix = path.rstrip("/") + "/"
        children = [p for p in self.tree if p.startswith(prefix) and "/" not in p[len(prefix) :]]
        body = "".join(response(p, self.tree[p]) for p in [path, *children])
        xml = f'<d:multistatus xmlns:d="DAV:" xmlns:oc="http://owncloud.org/ns" xmlns:nc="http://nextcloud.org/ns">{body}</d:multistatus>'
        return Reply(207, xml.encode())

    def get(self, url, timeout=None, stream=False):
        path = self._path(url)
        return Reply(200, self.files[path]) if path in self.files else Reply(404)


def nextcloud(tree=TREE, files=FILES) -> Nextcloud:
    source = Nextcloud(SHARE)
    source.http = FakeServer(tree, files)
    return source


def test_a_share_link_names_the_server_and_the_token():
    assert parse_share("https://cloud.example.org/s/Tok3n") == ("https://cloud.example.org", "Tok3n")
    assert parse_share("https://example.org/nextcloud/index.php/s/Tok3n/") == ("https://example.org/nextcloud", "Tok3n")
    for url in ["https://cloud.example.org/apps/files/", "https://cloud.example.org/s/", "ftp://x/s/Tok3n"]:
        with pytest.raises(ValueError):
            parse_share(url)


def test_the_share_token_and_password_are_the_login():
    assert Nextcloud(SHARE, "secret").http.auth == ("Tok3n", "secret")


def test_the_tree_is_listed_with_word_files_and_folders():
    entries, folders = nextcloud().walk()
    assert [(f.path, f.id) for f in folders] == [(["1 Work"], "10"), (["1 Work", "Sub folder"], "13")]
    by_name = {e.name: e for e in entries}
    assert set(by_name) == {"A post.docx", "Second.docx"}  # an .odt is not read
    post, second = by_name["A post.docx"], by_name["Second.docx"]
    assert (post.id, post.path, post.stem, post.kind, post.format) == ("11", ["1 Work"], "A post", "doc", "word")
    assert post.modified == "2026-10-20T08:00:00Z"
    assert post.created == "2026-01-01T00:00:00Z"  # no creation time: the upload time
    assert second.created == "2026-01-02T00:00:00Z"
    assert second.path == ["1 Work", "Sub folder"]


def test_an_unknown_share_stops_the_sync_instead_of_emptying_the_site():
    with pytest.raises(requests.HTTPError):
        nextcloud({}).walk()


def test_links_go_to_the_share_page():
    source = nextcloud()
    entries, folders = source.walk()
    assert source.url == SHARE
    assert source.embed == ""
    assert source.folder_url(folders[1]) == f"{SHARE}?dir=%2F1+Work%2FSub+folder"
    assert source.file_url(entries[0]) == f"{SHARE}?dir=%2F1+Work"


def test_a_share_syncs_and_builds_into_a_site(tmp_path):
    settings = config.Settings(SHARE, {}, tmp_path, "https://example.org", "Ann", False)
    assert sync.sync(settings, nextcloud()) is True
    content = cache.read_json(settings.cache_dir / "content" / "11.json")
    assert content["title"] == "A title"
    assert "First words." in content["html"]
    render.build(settings)
    assert "First words." in (settings.public_dir / "work" / "a-title" / "index.html").read_text()


def test_an_opendocument_file_is_skipped_with_a_warning_once(caplog):
    warned.clear()
    with caplog.at_level(logging.WARNING, logger="foldocs.sources.nextcloud"):
        nextcloud().walk()
        nextcloud().walk()  # the next sync, a minute later in `loop`
    assert [r.getMessage() for r in caplog.records] == [
        "Skipping 1 Work/notes.odt: an OpenDocument file; save it as Word (.docx) to publish it"
    ]


def test_a_missing_or_garbled_date_never_stops_the_sync():
    tree = {
        "/": {},
        "/Bad.docx": {"fileid": "1", "type": DOCX, "modified": "not a date", "creation_time": "", "upload_time": "x"},
        "/Upload only.docx": {"fileid": "2", "type": DOCX, "modified": "garbled", "upload_time": "1767225600"},
        "/Overflow.docx": {"fileid": "3", "type": DOCX, "creation_time": "99999999999999999999"},
    }
    entries = {e.id: e for e in nextcloud(tree).walk()[0]}
    assert (entries["1"].created, entries["1"].modified) == ("1970-01-01T00:00:00Z", "1970-01-01T00:00:00Z")
    assert (entries["2"].created, entries["2"].modified) == ("2026-01-01T00:00:00Z", "2026-01-01T00:00:00Z")
    assert entries["3"].created == entries["3"].modified == "2026-10-19T10:00:00Z"  # the default in response()


def test_the_default_themes_drive_pages_are_left_out_for_a_share(tmp_path):
    """The demo and how-it-works pages explain a Drive folder; a share's site goes without them."""
    settings = config.Settings(SHARE, {}, tmp_path, "https://example.org", "Ann", False)
    sync.sync(settings, nextcloud())
    render.build(settings)
    public = settings.public_dir
    assert not (public / "demo").exists() and not (public / "how-it-works").exists()
    assert "/demo/" not in (public / "sitemap.xml").read_text()
    drive = replace(settings, folder_url="f", output="drive")
    render.build(drive)
    assert (drive.public_dir / "demo" / "index.html").exists()


def test_a_file_id_that_is_not_a_number_is_not_trusted(tmp_path):
    """The ID names the file's cache entry: from a hostile server, "../x" would write outside it."""
    tree = {
        "/": {},
        "/1 Work": {"fileid": "../../folder"},
        "/1 Work/A post.docx": {**TREE["/1 Work/A post.docx"], "fileid": "../../../escaped"},
    }
    entries, folders = nextcloud(tree).walk()
    assert all(len(x.id) == 24 and x.id.isalnum() for x in [*entries, *folders])
    settings = config.Settings("", {}, tmp_path / "data", "https://example.org", "Ann", False)
    sync.sync(settings, nextcloud(tree))
    assert sorted(p.name for p in tmp_path.iterdir()) == ["data"]  # nothing written beside it


def test_the_folder_itself_is_recognized_behind_a_proxy_that_rewrites_the_prefix():
    from foldocs.sources.nextcloud import _in_share

    assert _in_share("/cloud/public.php/webdav/My%20site/") == _in_share("/public.php/webdav/My site") == "/My site"


@pytest.mark.parametrize(
    "answer",
    [
        # Entities that expand a billionfold, or nest deep enough to overflow the stack (CVE-2024-8176).
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        b'<d:multistatus xmlns:d="DAV:">&b;</d:multistatus>',
        # A DTD alone, with nothing declared, is refused too: an answer never has one.
        b'<?xml version="1.0"?><!DOCTYPE d><d:multistatus xmlns:d="DAV:"/>',
    ],
)
def test_an_answer_with_a_dtd_is_refused_before_it_is_expanded(answer):
    source = Nextcloud(SHARE)
    source.http = FakeServer()
    source.http.request = lambda *a, **k: Reply(207, answer)
    with pytest.raises(ET.ParseError):
        source.walk()


def test_a_share_listing_more_than_max_items_is_refused():
    source = nextcloud()
    source.max_items = 2
    with pytest.raises(TooLarge):
        source.walk()


def test_a_listing_over_max_bytes_is_refused():
    source = nextcloud()
    source.max_bytes = 100
    with pytest.raises(TooLarge):
        source.walk()


def jpeg() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "teal").save(buf, "JPEG")
    return buf.getvalue()


# A Markdown file as Nextcloud's Text app saves it: a pasted image goes into a hidden folder beside it.
TEXT_TREE = {
    "/": {},
    "/2 Notes": {"fileid": "20"},
    "/2 Notes/Trip.md": {"fileid": "21", "type": "text/markdown"},
    "/2 Notes/.attachments.21": {"fileid": "22"},
    "/2 Notes/.attachments.21/beach day.jpg": {"fileid": "23", "type": "image/jpeg"},
    "/2 Notes/notes.txt": {"fileid": "24", "type": "text/plain"},
}
TEXT_FILES = {
    "/2 Notes/Trip.md": b"# A trip\n\nSand.\n\n![The beach](.attachments.21/beach%20day.jpg)\n",
    "/2 Notes/.attachments.21/beach day.jpg": jpeg(),
}


def test_a_markdown_file_and_its_images_from_nextcloud_text(tmp_path):
    """Its images are on its page, but no section or photo of their own."""
    entries, _ = nextcloud(TEXT_TREE, TEXT_FILES).walk()
    assert {e.name for e in entries} == {"Trip.md", "beach day.jpg"}  # not the .txt

    settings = config.Settings(SHARE, {}, tmp_path, "https://example.org", "Ann", False)
    sync.sync(settings, nextcloud(TEXT_TREE, TEXT_FILES))
    assert {e["id"] for e in cache.read_json(settings.cache_dir / "manifest.json")} == {"21"}
    assert [f["path"] for f in cache.read_json(settings.cache_dir / "folders.json")] == [["2 Notes"]]
    content = cache.read_json(settings.cache_dir / "content" / "21.json")
    assert content["title"] == "A trip" and content["cover"]
    render.build(settings)
    page = (settings.public_dir / "notes" / "a-trip" / "index.html").read_text()
    assert "Sand." in page and 'alt="The beach"' in page
    assert not any("attachments" in p.name for p in settings.public_dir.rglob("*"))


def test_a_link_to_a_file_on_the_same_server_names_its_id():
    source = Nextcloud(SHARE)
    assert source.file_id("https://cloud.example.org/f/123") == "123"
    assert source.file_id("https://cloud.example.org/index.php/f/45") == "45"
    assert source.file_id("https://other.example.org/f/1") is None
    assert source.file_id(SHARE) is None
