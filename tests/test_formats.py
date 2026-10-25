"""What every format must offer, and how a file is told to be one: a new format passes this
before anything else (see foldocs/formats on adding one)."""

import pytest

from foldocs import convert, formats
from foldocs.formats import DOC, DOCX, Document, Reading, WordFile, format_for
from foldocs.sources import Entry
from foldocs.sources.drive import Drive
from foldocs.sources.nextcloud import Nextcloud


@pytest.mark.parametrize("fmt", formats.FORMATS, ids=lambda f: f.name)
def test_every_format_says_what_it_is(fmt):
    assert fmt.kind in ("doc", "image") and isinstance(fmt, Document) == (fmt.kind == "doc")
    assert fmt.name and fmt.label and fmt.mimes and formats.named(fmt.name) is fmt
    assert all(ext.startswith(".") and ext == ext.lower() for ext in fmt.extensions)
    for mime in fmt.mimes:
        assert format_for(mime, "any name") is fmt  # no two formats claim a type


class FakeSource:
    kind = "fake"

    def __init__(self, files):
        self.files = files

    def download(self, entry):
        return self.files[entry.name]


def word_file(text):
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("A title", style="Title")
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_a_documents_format_reads_it_into_clean_html_for_convert():
    entry = Entry("a", "a.docx", DOCX, [], "", "")
    text = entry.reader.read(Reading(entry, FakeSource({"a.docx": word_file("Some words.")})))
    doc = convert.convert(text.html, lambda src: None, lambda info, sizes=None: {}, text.title, text.subtitle)
    assert (doc.title, doc.html) == ("A title", "<p>Some words.</p>")


def test_a_google_doc_is_only_read_from_a_service_that_exports_it():
    entry = Entry("a", "Notes", DOC, [], "", "")
    with pytest.raises(ValueError, match="no Google Docs"):
        entry.reader.read(Reading(entry, FakeSource({})))
    from foldocs.sources import Exports

    assert isinstance(Drive("key", "folder"), Exports)
    assert not isinstance(Nextcloud("https://cloud.example.org/s/Tok3n"), Exports)


@pytest.mark.parametrize(
    "mime, name, expected",
    [
        (DOC, "Notes", "gdoc"),
        (DOCX, "Notes.docx", "word"),
        ("application/octet-stream", "NOTES.DOCX", "word"),  # a server that doesn't know the type
        ("text/markdown", "Notes.md", "markdown"),
        ("text/x-markdown", "Notes.md", "markdown"),
        ("text/plain", "Notes.markdown", "markdown"),
        ("image/jpeg", "me.jpg", "image"),
        ("text/plain", "notes.txt", None),
        ("application/vnd.oasis.opendocument.text", "notes.odt", None),
        ("image/svg+xml", "logo.svg", None),
        ("application/pdf", "Notes.docx", None),  # a type that says otherwise wins
    ],
)
def test_a_file_is_known_by_its_type_or_else_its_extension(mime, name, expected):
    found = format_for(mime, name)
    assert (found.name if found else None) == expected


def test_an_entry_records_its_format_once():
    entry = Entry("a", "A post.docx", DOCX, [], "", "")
    assert (entry.format, entry.kind, entry.stem) == ("word", "doc", "A post")
    assert WordFile().stem("v1.0") == "v1.0"
    assert Entry("b", "intro.MD", "text/markdown", [], "", "").stem == "intro"
    with pytest.raises(ValueError):
        Entry("a", "notes.txt", "text/plain", [], "", "")


def test_the_page_is_made_from_clean_tags_only():
    """Whatever a format leaves in, styles and unknown tags never reach the site: only their text."""
    doc = convert.convert(
        '<p style="color:red" class="x"><span style="font-weight:700">a</span> <b>b</b> <font>c</font>'
        "<script>d</script></p><section><p>e</p></section>",
        lambda src: None,
        lambda info, sizes=None: {},
    )
    assert doc.html == "<p>a <strong>b</strong> cd</p>\n<p>e</p>"


def test_footnotes_lead_to_their_notes_and_back():
    """As mammoth writes a Word file's: the note's id is on its list item."""
    doc = convert.convert(
        '<p>Text<sup><a href="#footnote-1" id="footnote-ref-1">[1]</a></sup></p>'
        '<ol><li id="footnote-1"><p>A note. <a href="#footnote-ref-1">↑</a></p></li></ol>',
        lambda src: None,
        lambda info, sizes=None: {},
    )
    assert doc.html == (
        '<p>Text<sup><a href="#footnote-1" id="footnote-ref-1">[1]</a></sup></p>\n'
        '<ol><li id="footnote-1">A note. <a href="#footnote-ref-1">↑</a></li></ol>'
    )


def test_formats_are_named_for_people():
    assert formats.described() == "a Google Doc, a Word file, a Markdown file or an image"


def test_a_link_from_a_document_leads_to_a_file_by_its_path():
    doc = Entry("d", "post.md", "text/markdown", ["writing", "2026"], "", "")
    photo = Entry("p", "a b.jpg", "image/jpeg", ["writing", "photos"], "", "")
    attached = Entry("q", "x.png", "image/png", ["writing", "2026", ".attachments.12"], "", "")
    file = Reading(doc, FakeSource({}), {(*e.path, e.name): e for e in (doc, photo, attached)})
    assert file.linked("../photos/a%20b.jpg") is photo
    assert file.linked("./.attachments.12/x.png") is attached
    for src in ["../../../a b.jpg", "missing.jpg", "/writing/photos/a%20b.jpg", "https://example.org/x.png", "#top"]:
        assert file.linked(src) is None, src


def test_a_hidden_file_or_folder_starts_with_a_dot():
    assert Entry("a", "x.jpg", "image/jpeg", ["Trip", ".attachments.12"], "", "").hidden
    assert Entry("a", ".x.jpg", "image/jpeg", [], "", "").hidden
    assert not Entry("a", "x.jpg", "image/jpeg", ["Trip"], "", "").hidden
